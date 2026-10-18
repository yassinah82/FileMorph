import os
import asyncio
from pathlib import Path
from typing import List, Optional, Callable, Dict, Any
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches

try:
    import winocr
    HAS_WINOCR = True
except ImportError:
    HAS_WINOCR = False

class OCREngine:
    """
    High-accuracy OCR Engine leveraging native Windows Media OCR (Windows 10/11 WinRT).
    Zero external third-party binary installations required!
    """

    @staticmethod
    def is_available() -> bool:
        return HAS_WINOCR

    @staticmethod
    def get_supported_languages() -> List[str]:
        """Returns list of language tags installed in Windows (e.g. ['en-US', 'ar-SA'])."""
        if not HAS_WINOCR:
            return []
        try:
            langs = winocr.installed_recognizer_languages()
            return [l.language_tag for l in langs] if langs else ["en-US"]
        except Exception:
            return ["en-US"]

    @classmethod
    def recognize_pil_image(cls, img: Image.Image, lang: str = "en-US") -> str:
        """Runs Windows Native OCR on a PIL Image."""
        if not HAS_WINOCR:
            raise RuntimeError("winocr is not installed. Please install winocr.")
        
        # WinOCR requires RGB or RGBA mode
        if img.mode not in ("RGB", "RGBA"):
            img = img.convert("RGB")

        # winocr is async, run in event loop
        loop = asyncio.new_event_loop()
        try:
            asyncio.set_event_loop(loop)
            result = loop.run_until_complete(winocr.recognize_pil(img, lang))
            return result.text if hasattr(result, "text") else str(result)
        finally:
            loop.close()

    @classmethod
    def recognize_image_file(cls, image_path: str, lang: str = "en-US") -> str:
        """Extracts text from an image file."""
        src = Path(image_path).resolve()
        if not src.exists():
            raise FileNotFoundError(f"Image not found: {src}")

        with Image.open(src) as img:
            return cls.recognize_pil_image(img, lang=lang)

    @classmethod
    def recognize_pdf_pages(
        cls,
        pdf_path: str,
        lang: str = "en-US",
        dpi: int = 200,
        progress_callback: Optional[Callable[[float, str], None]] = None
    ) -> List[str]:
        """Renders PDF pages and extracts text using Windows OCR."""
        import pymupdf as fitz
        src = Path(pdf_path).resolve()
        doc = fitz.open(str(src))
        total_pages = len(doc)
        pages_text: List[str] = []

        zoom = dpi / 72.0
        matrix = fitz.Matrix(zoom, zoom)

        try:
            for idx in range(total_pages):
                if progress_callback:
                    pct = (idx / total_pages) * 0.9
                    progress_callback(pct, f"OCR scanning page {idx + 1} of {total_pages}...")

                page = doc.load_page(idx)
                pix = page.get_pixmap(matrix=matrix, alpha=False)
                pil_img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)

                page_text = cls.recognize_pil_image(pil_img, lang=lang)
                pages_text.append(page_text)

            if progress_callback:
                progress_callback(1.0, f"OCR completed for {total_pages} pages.")
        finally:
            doc.close()

        return pages_text

    @classmethod
    def to_docx_ocr(
        cls,
        source_path: str,
        output_path: Optional[str] = None,
        lang: str = "en-US",
        progress_callback: Optional[Callable[[float, str], None]] = None
    ) -> str:
        """
        Converts an image or scanned PDF into an editable Word (.docx) document via OCR.
        """
        src = Path(source_path).resolve()
        if not output_path:
            output_path = str(src.with_suffix(".docx"))
        else:
            output_path = str(Path(output_path).resolve())

        if progress_callback:
            progress_callback(0.1, "Initializing OCR pipeline...")

        doc = Document()
        # Set clean document styling
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)

        src_ext = src.suffix.lower()

        if src_ext == ".pdf":
            pages_text = cls.recognize_pdf_pages(str(src), lang=lang, progress_callback=progress_callback)
            for idx, text in enumerate(pages_text):
                if idx > 0:
                    doc.add_page_break()
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.line_spacing = 1.15
                doc.add_paragraph(text)
        else:
            if progress_callback:
                progress_callback(0.4, "Running OCR on image...")
            text = cls.recognize_image_file(str(src), lang=lang)
            if progress_callback:
                progress_callback(0.8, "Creating Word document...")
            doc.add_paragraph(text)

        doc.save(output_path)

        if progress_callback:
            progress_callback(1.0, "Saved OCR Word document.")

        return output_path
